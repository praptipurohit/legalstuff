"""
Legal Contracts MIS — Excel Builder
Generates LegalMIS_v2.xlsx with all sheets, headers, formatting, and dropdowns.
After opening in Excel: File > Save As > .xlsm, then import the .bas VBA files.
"""

import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

# ── Styles ──────────────────────────────────────────────────────────────────
def fill(hex_color):
    return PatternFill("solid", fgColor=hex_color)

HDR_FILL       = fill("4472C4")
TITLE_DARK     = fill("1F3864")
STAGE_HDR_FILL = fill("404040")
HINT_FILL      = fill("2E74B5")
FORM_LABEL     = fill("EEF2F9")
GRAY_FILL      = fill("F9F9F9")
WHITE_FILL     = fill("FFFFFF")
ORANGE_ACCENT  = fill("ED7D31")

HDR_FONT        = Font(bold=True, color="FFFFFF", name="Calibri", size=10)
TITLE_FONT      = Font(bold=True, color="FFFFFF", name="Calibri", size=13)
STAGE_HDR_FONT  = Font(bold=True, color="FFFFFF", name="Calibri", size=9)
LABEL_FONT      = Font(bold=True, color="1F3864", name="Calibri", size=10)
BODY_FONT       = Font(name="Calibri", size=10)
SMALL_FONT      = Font(italic=True, color="595959", name="Calibri", size=8)
RED_ITALIC      = Font(italic=True, color="C00000", name="Calibri", size=8)

CENTER  = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT    = Alignment(horizontal="left",   vertical="center", wrap_text=False)
LWRAP   = Alignment(horizontal="left",   vertical="center", wrap_text=True)

def border(color="BFBFBF"):
    s = Side(style="thin", color=color)
    return Border(left=s, right=s, top=s, bottom=s)

def hdr_border():
    s = Side(style="medium", color="FFFFFF")
    return Border(left=s, right=s, top=s, bottom=s)

# ── Reference data ───────────────────────────────────────────────────────────
STATUS_CSV = (
    "Executed / Closed,Reverted to Counter-party,Pending with Counter-party,"
    "Under Review,Pending Execution,Pending with Business,"
    "Not Responded,On Hold,Abandoned,N/A - Advisory"
)
DIVISION_CSV = "Enterprise,Telecom,Media,Digital,Retail,Corporate,Other"
# Edit these reviewer names to match your team
REVIEWER_CSV = "Neel,Faizan Khan,Kulwinder Singh,Priya K,Amit S,Kulwinder,Praptipurohit"

MAIN_HEADERS = [
    "S. No.", "Document Name", "Vendor / Party", "Business SPOC",
    "Business Division", "First Reviewer", "Second Reviewer",
    "Date of Receipt", "Date of Reply", "Contract Value", "Status"
]
MAIN_WIDTHS = [8, 38, 30, 22, 18, 18, 18, 16, 16, 18, 30]

# ── Helpers ──────────────────────────────────────────────────────────────────
def write_header_row(ws, headers, widths=None, row=1, start_col=1):
    for i, h in enumerate(headers):
        c = ws.cell(row=row, column=start_col + i, value=h)
        c.fill = HDR_FILL
        c.font = HDR_FONT
        c.alignment = CENTER
        c.border = hdr_border()
    if widths:
        for i, w in enumerate(widths):
            ws.column_dimensions[get_column_letter(start_col + i)].width = w
    ws.row_dimensions[row].height = 22

def add_dropdown(ws, col_letter, csv_values, start_row=2, end_row=2000,
                 error_msg=True, err_title="Invalid Value",
                 err_msg="Please select from the list."):
    dv = DataValidation(
        type="list",
        formula1=f'"{csv_values}"',
        allow_blank=True,
        showErrorMessage=error_msg,
        errorTitle=err_title,
        error=err_msg
    )
    dv.sqref = f"{col_letter}{start_row}:{col_letter}{end_row}"
    ws.add_data_validation(dv)

def style_cell(cell, value=None, fnt=None, fll=None, aln=None, brd=None, fmt=None):
    if value is not None: cell.value = value
    if fnt:  cell.font      = fnt
    if fll:  cell.fill      = fll
    if aln:  cell.alignment = aln
    if brd:  cell.border    = brd
    if fmt:  cell.number_format = fmt

# ── Sheet: Master ────────────────────────────────────────────────────────────
def build_master(wb):
    ws = wb.create_sheet("Master")
    write_header_row(ws, MAIN_HEADERS, MAIN_WIDTHS)
    ws.freeze_panes = "A2"

    add_dropdown(ws, "E", DIVISION_CSV, error_msg=False)
    add_dropdown(ws, "F", REVIEWER_CSV, error_msg=False)
    add_dropdown(ws, "G", REVIEWER_CSV, error_msg=False)
    add_dropdown(ws, "K", STATUS_CSV,
                 error_msg=True,
                 err_title="Invalid Status",
                 err_msg="Please select a valid status from the list.")

    for row in range(2, 2001):
        ws.cell(row=row, column=8).number_format = "DD-MMM-YYYY"
        ws.cell(row=row, column=9).number_format = "DD-MMM-YYYY"

    ws.sheet_view.showGridLines = True
    return ws

# ── Sheet: Quick Entry ───────────────────────────────────────────────────────
def build_quick_entry(wb):
    ws = wb.create_sheet("Quick Entry")
    ws.sheet_view.showGridLines = False

    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 26
    ws.column_dimensions["C"].width = 42
    ws.column_dimensions["D"].width = 4

    # Title row
    ws.merge_cells("B1:C1")
    style_cell(ws["B1"],
               value="  NEW CONTRACT — QUICK ENTRY",
               fnt=TITLE_FONT, fll=TITLE_DARK, aln=CENTER)
    ws.row_dimensions[1].height = 30

    # Shortcut hints
    ws.merge_cells("B2:C2")
    style_cell(ws["B2"],
               value=("Ctrl+Shift+S  Stage entry    │    "
                      "Ctrl+Shift+A  Commit to Master    │    "
                      "Ctrl+Shift+F  Clear form    │    "
                      "Ctrl+Shift+X  Clear staging"),
               fnt=SMALL_FONT, fll=fill("E9EDF4"), aln=CENTER)
    ws.row_dimensions[2].height = 14

    # Blank spacer
    ws.row_dimensions[3].height = 6

    # Form fields
    FIELDS = [
        (4,  "Document Name",      False, None,         False),
        (5,  "Vendor / Party",     False, None,         False),
        (6,  "Business SPOC",      False, None,         False),
        (7,  "Business Division",  True,  DIVISION_CSV, False),
        (8,  "First Reviewer",     True,  REVIEWER_CSV, False),
        (9,  "Second Reviewer",    True,  REVIEWER_CSV, False),
        (10, "Date of Receipt",    False, None,         True),
        (11, "Date of Reply",      False, None,         True),
        (12, "Contract Value",     False, None,         False),
        (13, "Status",             True,  STATUS_CSV,   False),
    ]

    for row, label, has_dv, csv, is_date in FIELDS:
        # Label
        lc = ws.cell(row=row, column=2, value=f"  {label}")
        lc.font = LABEL_FONT
        lc.fill = FORM_LABEL
        lc.alignment = LEFT
        lc.border = border("C8C8C8")

        # Input
        ic = ws.cell(row=row, column=3)
        ic.fill = WHITE_FILL
        ic.font = BODY_FONT
        ic.alignment = LEFT
        ic.border = border("4472C4")
        if is_date:
            ic.number_format = "DD-MMM-YYYY"
        if has_dv and csv:
            dv = DataValidation(
                type="list",
                formula1=f'"{csv}"',
                allow_blank=True,
                showErrorMessage=(row == 13),
                errorTitle="Invalid Value",
                error="Please select from the dropdown."
            )
            dv.sqref = f"C{row}"
            ws.add_data_validation(dv)
        ws.row_dimensions[row].height = 24

    # Required note
    ws.merge_cells("B14:C14")
    style_cell(ws["B14"],
               value="  Document Name, Vendor and Status are required before staging.",
               fnt=RED_ITALIC, fll=fill("FFF2F2"), aln=LEFT)
    ws.row_dimensions[14].height = 14

    # Staging banner
    ws.merge_cells("B15:C15")
    style_cell(ws["B15"],
               value="▼   STAGING AREA  —  Entries queued for commit to Master   ▼",
               fnt=Font(bold=True, color="FFFFFF", name="Calibri", size=10),
               fll=HINT_FILL, aln=CENTER)
    ws.row_dimensions[15].height = 22

    # Staging headers (row 16, cols A-K)
    for i, h in enumerate(MAIN_HEADERS, 1):
        c = ws.cell(row=16, column=i, value=h)
        c.fill = STAGE_HDR_FILL
        c.font = STAGE_HDR_FONT
        c.alignment = CENTER
        c.border = border("808080")
    ws.row_dimensions[16].height = 20

    # Staging data rows 17-46
    for row in range(17, 47):
        bg = fill("F7F9FF") if row % 2 == 0 else WHITE_FILL
        for col in range(1, 12):
            c = ws.cell(row=row, column=col)
            c.fill = bg
            c.font = BODY_FONT
            c.alignment = LEFT
            c.border = border("D0D0D0")
        ws.row_dimensions[row].height = 20

    ws.freeze_panes = "A17"
    return ws

# ── Sheet: Paste_NewWork ─────────────────────────────────────────────────────
def build_paste_sheet(wb, name):
    ws = wb.create_sheet(name)
    headers = MAIN_HEADERS + ["Duplicate Check"]
    widths  = MAIN_WIDTHS  + [20]
    write_header_row(ws, headers, widths)

    # Override Duplicate Check header colour
    dc = ws.cell(row=1, column=12)
    dc.fill = ORANGE_ACCENT
    dc.font = HDR_FONT

    ws.freeze_panes = "A2"
    add_dropdown(ws, "E", DIVISION_CSV, error_msg=False)
    add_dropdown(ws, "F", REVIEWER_CSV, error_msg=False)
    add_dropdown(ws, "G", REVIEWER_CSV, error_msg=False)
    add_dropdown(ws, "K", STATUS_CSV,
                 error_msg=True, err_title="Invalid Status",
                 err_msg="Please select a valid status.")

    for row in range(2, 2001):
        ws.cell(row=row, column=8).number_format = "DD-MMM-YYYY"
        ws.cell(row=row, column=9).number_format = "DD-MMM-YYYY"

    # Instructions in row 2 comment area (lightly styled)
    ws.merge_cells("A2:K2")
    style_cell(ws["A2"],
               value=("  After pasting reviewer data here, run:  "
                      "Alt+F8 → CheckForDuplicates → Run"),
               fnt=SMALL_FONT, fll=fill("FFF9F0"), aln=LEFT)
    ws.row_dimensions[2].height = 16
    return ws

# ── Sheet: Paste_StatusUpdates ───────────────────────────────────────────────
def build_paste_status(wb):
    ws = build_paste_sheet(wb, "Paste_StatusUpdates")
    ws.merge_cells("A2:K2")
    style_cell(ws["A2"],
               value=("  After pasting status updates here, run:  "
                      "Alt+F8 → ReconcileStatusUpdates → Run"),
               fnt=SMALL_FONT, fll=fill("FFF9F0"), aln=LEFT)
    return ws

# ── Sheet: StatusReconciliation ──────────────────────────────────────────────
def build_reconciliation(wb):
    ws = wb.create_sheet("StatusReconciliation")
    headers = ["Master Row", "Document Name", "Vendor / Party",
               "Status in Master", "Proposed Status", "Changed?", "Action"]
    widths  = [12, 38, 30, 28, 28, 12, 22]
    write_header_row(ws, headers, widths)
    ws.freeze_panes = "A2"

    ws.merge_cells("A2:G2")
    style_cell(ws["A2"],
               value=("  Run: Alt+F8 → ReconcileStatusUpdates   "
                      "  Then review Action column (change 'Pending' → 'Skip' to ignore)   "
                      "  Then: Alt+F8 → ApplyApprovedUpdates"),
               fnt=SMALL_FONT, fll=fill("EEF2F9"), aln=LEFT)
    ws.row_dimensions[2].height = 16
    return ws

# ── Sheet: MonthlySummary ────────────────────────────────────────────────────
def build_monthly_summary(wb):
    ws = wb.create_sheet("MonthlySummary")
    ws.sheet_view.showGridLines = False

    ws.column_dimensions["A"].width = 36
    ws.column_dimensions["B"].width = 12
    ws.column_dimensions["C"].width = 14
    ws.column_dimensions["D"].width = 18
    ws.column_dimensions["E"].width = 16

    ws.merge_cells("A1:E1")
    style_cell(ws["A1"],
               value="LEGAL CONTRACTS MIS — MONTHLY SUMMARY",
               fnt=Font(bold=True, color="FFFFFF", name="Calibri", size=14),
               fll=TITLE_DARK, aln=CENTER)
    ws.row_dimensions[1].height = 32

    ws.merge_cells("A2:E2")
    style_cell(ws["A2"],
               value="⚡  Refresh: Alt+F8 → RefreshMonthlySummary → Run",
               fnt=SMALL_FONT, fll=fill("E9EDF4"), aln=CENTER)
    ws.row_dimensions[2].height = 16

    # Status legend table
    STATUS_COLORS_HEX = {
        "Executed / Closed":          "C6EFCE",
        "N/A - Advisory":             "C6EFCE",
        "Reverted to Counter-party":  "FFC783",
        "Pending with Counter-party": "FFC783",
        "Pending Execution":          "FFC783",
        "Pending with Business":      "FFC783",
        "Under Review":               "FFC7C7",
        "Not Responded":              "FFC7C7",
        "On Hold":                    "BDD7EE",
        "Abandoned":                  "BDD7EE",
    }

    ws.row_dimensions[3].height = 8

    ws.merge_cells("A4:E4")
    style_cell(ws["A4"], value="STATUS COLOUR LEGEND",
               fnt=Font(bold=True, color="FFFFFF", name="Calibri", size=10),
               fll=fill("2E74B5"), aln=CENTER)
    ws.row_dimensions[4].height = 20

    for i, (status, hex_c) in enumerate(STATUS_COLORS_HEX.items(), 5):
        c = ws.cell(row=i, column=1, value=status)
        c.fill = fill(hex_c)
        c.font = BODY_FONT
        c.alignment = LEFT
        c.border = border()
        ws.row_dimensions[i].height = 18

    return ws

# ── Build ────────────────────────────────────────────────────────────────────
def build_all():
    wb = openpyxl.Workbook()
    wb.remove(wb.active)        # remove default Sheet

    build_master(wb)
    build_quick_entry(wb)
    build_paste_sheet(wb, "Paste_NewWork")
    build_paste_status(wb)
    build_reconciliation(wb)
    build_monthly_summary(wb)

    out = "/home/user/legalstuff/LegalMIS_v2.xlsx"
    wb.save(out)
    print(f"✓  Saved: {out}")
    return out

if __name__ == "__main__":
    build_all()
